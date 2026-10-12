"""Build a throwaway Shadow Filer for clicking through the console without touching real files.

It copies the code and the built UI into <target>/agent (so its search.db, cleanup report and logs are its
own: they live next to the code), creates made-up inbox and PARA folders with fictional files, some carrying
a "downloaded from" mark, plus made-up firmware, document versions, pictures and a program for the Library
page, and writes a config.yaml on port 8766.

    python tests/make_test_agent.py <target folder> [--tiny-bin] [--proven]
    cd <target folder>/agent
    <venv python> shadow_agent.py --config config.yaml run        -> http://127.0.0.1:8766

--tiny-bin makes the copy pretend the Recycle Bin holds 1 KB, to exercise "too big, delete permanently".
--proven records 20 agreed Inbox decisions for two folders (Payroll by same website, Acme-Rollout by file
name), so auto-filing can be watched: switch it on in Settings, wait a minute, see the move, undo it.
Files it recycles land in the real Recycle Bin; Show and Open launch Explorer / an app on the desktop.
Stop it with Ctrl+C (or end the process on port 8766) and delete the target folder afterwards.
"""
import os
import shutil
import sys
import textwrap
import time
from pathlib import Path

SRC = Path(__file__).resolve().parent.parent
args = [a for a in sys.argv[1:] if not a.startswith("--")]
if len(args) != 1:
    sys.exit(__doc__)
T = Path(args[0]).resolve()
if T == SRC or SRC in T.parents:
    sys.exit("Choose a target outside the shadow-filer folder.")
if T.exists():
    shutil.rmtree(T)
agent, inbox, para = T / "agent", T / "inbox", T / "para"
agent.mkdir(parents=True)
for f in list(SRC.glob("*.py")) + [SRC / "ui.html"]:
    shutil.copy2(f, agent / f.name)
shutil.copytree(SRC / "ui_dist", agent / "ui_dist")

if "--tiny-bin" in sys.argv:  # only in the copy: pretend the Recycle Bin holds at most 1 KB
    code = (agent / "cleanup.py").read_text(encoding="utf-8")
    needle = "def recycle_bin_limit(path: str) -> dict:\n"
    assert code.count(needle) == 1
    code = code.replace(needle, needle + '    return {"known": True, "nuke": False, "max_bytes": 1024, '
                                         '"drive": "TEST", "volume_bytes": 0}\n')
    (agent / "cleanup.py").write_text(code, encoding="utf-8")


def w(p: Path, text: str, site: str | None = None):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(textwrap.dedent(text).strip() + "\n", encoding="utf-8")
    if site and os.name == "nt":  # the mark browsers leave on downloads (origin.py reads it)
        with open(str(p) + ":Zone.Identifier", "w", encoding="utf-8", newline="") as f:
            f.write(f"[ZoneTransfer]\r\nZoneId=3\r\nReferrerUrl=https://{site}/\r\nHostUrl=https://{site}/download/file\r\n")


topics = {
    "1-Projects/Acme-Rollout": ("Acme rollout", "The Acme network rollout covers site surveys, switch staging, "
                                "cutover rehearsal and sign-off. Milestone {i}: wave {i} migrated.", "mail.example.com"),
    "2-Areas/Finance": ("Invoice", "Tax invoice INV-2026-{i:03d} from Northwind Supplies. Amount due 4,{i}50.00, "
                        "net 30, GST included.", "mail.example.com"),
    "2-Areas/Payroll": ("Payslip", "Payslip for month {i}: basic pay, allowances, deductions, net pay credited "
                        "to bank account.", "portal.payroll.example"),
    "3-Resources/Manuals": ("Manual", "User manual chapter {i}: installing the router, firmware upgrade, factory "
                            "reset, LED status table.", None),
}
for folder, (title, body, site) in topics.items():
    for i in range(1, 7):
        w(para / folder / f"{title.replace(' ', '_')}_{i}.txt", f"{title} {i}\n\n" + body.format(i=i), site)
w(inbox / "statement-october.txt", "Monthly statement: pay, allowances, deductions and net pay for October.",
  "portal.payroll.example")
w(inbox / "Acme_status_report.txt", "Acme rollout weekly status. Wave 12 migrated, cutover rehearsal passed.",
  "mail.example.com")
w(inbox / "GR-One kickoff notes.txt", "Kickoff notes for the new GR One engagement: stakeholders and timeline.")
w(inbox / "sf_open_test_9f3.txt", "File used to test the Show and Open buttons.")
w(inbox / "big_report.txt", "Quarterly report draft, version A. " * 400)
(inbox / "big_report (1).txt").write_bytes((inbox / "big_report.txt").read_bytes())      # duplicate, over 1 KB
w(inbox / "small_note.txt", "tiny " * 20)
(inbox / "small_note (1).txt").write_bytes((inbox / "small_note.txt").read_bytes())      # duplicate, under 1 KB
# Library page: firmware and an unpacked appliance, versions of one document, look-alike pictures, a program
for name, n in (("FGT_100F-v7.4.8.M-build2795-FORTINET.out", 4000), ("FGT_100F-v7.6.2.M-build3462-FORTINET.out", 4200)):
    (inbox / name).write_bytes(b"\0" * n)
vm = inbox / "Northwind-NAC-9.2.1-411.ovf"
vm.mkdir()
(vm / "Northwind-NAC.ovf").write_bytes(b"\0" * 900)
(vm / "disk1.vmdk").write_bytes(b"\0" * 20000)
w(para / "1-Projects/Acme-Rollout" / "Acme rollout proposal.docx", "Proposal, first issue.")
w(inbox / "Acme rollout proposal_v2.docx", "Proposal, second issue, with the revised timeline.")
w(inbox / "Acme rollout proposal_v3_Final.docx", "Proposal, third and final issue, with the revised timeline and prices.")
(inbox / "lab-tool-setup.exe").write_bytes(b"MZ" + b"\0" * 2000)   # not a real program: Windows reports it unsigned
try:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (800, 600), "white")
    d = ImageDraw.Draw(im)
    for i in range(400):   # busy enough to be a few KB: tiny icons are not compared
        x, y = (i * 137) % 760, (i * 251) % 560
        d.ellipse([x, y, x + 20 + i % 60, y + 15 + i % 45], fill=((i * 18) % 255, 255 - (i * 15) % 255, (i * 60) % 255))
    im.save(inbox / "team-photo.png")
    im.resize((400, 300)).save(inbox / "team-photo-small.jpg")
    im.save(para / "3-Resources/Manuals" / "team-photo.png")
except ImportError:
    pass
# Preview panel: a real PDF (two pages), a locked PDF and a real Word file, so the document view has something to draw
try:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    def pdf(path, lines, password=None):
        wr = PdfWriter()
        for n, line in enumerate(lines, 1):
            page = wr.add_blank_page(width=595, height=842)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): wr._add_object(font)})})
            st = DecodedStreamObject()
            st.set_data(f"BT /F1 20 Tf 60 760 Td ({line}) Tj 0 -40 Td /F1 12 Tf (Page {n} of {len(lines)}) Tj ET".encode())
            page[NameObject("/Contents")] = wr._add_object(st)
        if password:
            wr.encrypt(user_password=password, owner_password=password + "-owner", algorithm="AES-256")
        with open(path, "wb") as f:
            wr.write(f)

    pdf(inbox / "Acme site survey.pdf", ["Acme rollout: site survey, Leeds branch", "Findings: cabling ready, rack space for two switches"])
    pdf(inbox / "Payroll summary (protected).pdf", ["Not for the agent"], password="made-up-password")
except ImportError:
    pass
try:
    import docx
    d = docx.Document()
    d.add_heading("Acme rollout: kickoff minutes", 1)
    d.add_paragraph("Attendees: project lead, site lead, network engineer. All names in this file are made up.")
    d.add_heading("Decisions", 2)
    for line in ("Wave 12 moves to the first week of November.", "Cutover rehearsal is repeated before each wave.",
                 "The weekly status report goes out on Fridays."):
        d.add_paragraph(line, style="List Bullet")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Action", "Owner", "Due")):
        c.text = h
    for row in (("Order switches for wave 12", "Site lead", "20 Oct"), ("Book the rehearsal window", "Project lead", "24 Oct")):
        for c, v in zip(t.add_row().cells, row):
            c.text = v
    d.save(str(inbox / "Acme kickoff minutes.docx"))
except ImportError:
    pass
old = time.time() - 40 * 86400   # Cleanup skips files changed in the last 10 minutes
for p in list(inbox.rglob("*")) + list(para.rglob("*")):
    if p.is_file():
        os.utime(p, (old, old))
for i, name in enumerate(("Acme rollout proposal_v2.docx", "Acme rollout proposal_v3_Final.docx")):
    os.utime(inbox / name, (old + (i + 1) * 86400,) * 2)   # each version a day newer than the filed one
(agent / "config.yaml").write_text(textwrap.dedent(f"""
    inbox_folders:
    - {inbox}
    para_root: {para}
    destination_depth: 2
    exclude_destinations: []
    ollama:
      url: http://localhost:11434
      model: qwen3:4b
      timeout_seconds: 120
      think: false
      startup_wait_seconds: 5
    extraction:
      max_chars: 4000
      max_pdf_pages: 5
    settle_seconds: 2
    hash_max_mb: 200
    reconcile_minutes: 30
    history_examples: 15
    catch_up_days: 0
    ui:
      enabled: true
      port: 8766
    database: shadow.db
    log_file: shadow.log
    ignore_patterns:
    - '*.tmp'
    rules: []
    cleanup:
      scan_hours: 0
    autofile:
      enabled: false
      delay_minutes: 1
      every_minutes: 1
    """).lstrip(), encoding="utf-8")

if "--proven" in sys.argv:  # the evidence 20 real "File it" clicks would leave in the journal
    sys.path.insert(0, str(agent))
    import mover
    from datetime import datetime, timedelta
    w(inbox / "payslip-november.txt", "Payslip for November: basic pay, allowances, deductions, net pay.",
      "portal.payroll.example")
    w(inbox / "Acme_cutover_plan.txt", "Acme rollout cutover plan: freeze window, rollback steps, sign-off.")
    c = mover._conn(agent / "shadow.db")
    for folder, kind, stem, site in (("2-Areas/Payroll", "origin", "Payslip_2025", "portal.payroll.example"),
                                     ("1-Projects/Acme-Rollout", "name", "Acme_wave", None)):
        for i in range(1, 21):
            dest = para / folder / f"{stem}_{i:02d}.txt"
            w(dest, f"{stem.replace('_', ' ')} {i}: filed earlier through the Inbox.", site)
            os.utime(dest, (old, old))
            c.execute("INSERT INTO moves (ts, src, dest, size, sha256, folder, source, rank, top_folder, top_score, "
                      "top_source, status) VALUES (?,?,?,?,?,?,?,1,?,NULL,?, 'moved')",
                      ((datetime.now() - timedelta(days=30 - i)).isoformat(timespec="seconds"), str(inbox / dest.name),
                       str(dest), dest.stat().st_size, mover.sha256_of(str(dest)), folder, kind, folder, kind))
    c.commit()
    c.close()
print(f"ready: {T}\nrun:   cd {agent}  &&  python shadow_agent.py --config config.yaml run   (http://127.0.0.1:8766)")
